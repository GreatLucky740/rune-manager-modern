import json
import openpyxl

with open("artifact_seed_swlens.json", encoding="utf-8") as f:
    seed = json.load(f)
wb = openpyxl.load_workbook(r"D:\Artifact tool 3.7.2022 -Public.xlsm", read_only=False, data_only=True)
ws = wb["JSON translate"]
units = {}
for row in range(1, ws.max_row + 1):
    mid, name = ws.cell(row, 4).value, ws.cell(row, 5).value
    if isinstance(mid, (int, float)) and isinstance(name, str) and name.strip():
        units[str(int(mid))] = name.strip()
payload = {
    "profiles": seed["profiles"], "rolls": seed["rolls"],
    "stat_names": seed["stat_names"], "unit_names": units,
    "config": {"keep": 6, "main_bonus": 0.75, "divisor": 2.5},
}
with open("outputs/artifact_manager_modern/Artifact_Manager_Data.json", "w", encoding="utf-8") as f:
    json.dump(payload, f, ensure_ascii=False, separators=(",", ":"))
print(len(payload["profiles"]), len(units))
