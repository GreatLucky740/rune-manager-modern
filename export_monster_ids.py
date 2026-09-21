import json
import openpyxl

src = r"D:\Artifact tool 3.7.2022 -Public.xlsm"
wb = openpyxl.load_workbook(src, read_only=False, data_only=True)
ws = wb["JSON translate"]
rows = []
for r in range(1, ws.max_row + 1):
    mid, name = ws.cell(r, 4).value, ws.cell(r, 5).value
    try:
        mid = int(mid)
    except (TypeError, ValueError):
        continue
    if isinstance(name, str) and name.strip():
        rows.append({"id": mid, "name": name.strip()})
rows = list({x["id"]: x for x in rows}.values())
with open("monster_ids.json", "w", encoding="utf-8") as f:
    json.dump(rows, f, ensure_ascii=False)
print(len(rows))
